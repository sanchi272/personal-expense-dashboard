
import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image

# Sample expense data
data = {
    "Date": [
        "01-10-2026", "03-10-2026", "05-10-2026",
        "08-10-2026", "12-10-2026", "15-10-2026",
        "18-10-2026", "20-10-2026", "22-10-2026",
        "25-10-2026"
    ],
    "Category": [
        "Food", "Travel", "Shopping", "Food", "Bills",
        "Travel", "Food", "Shopping", "Bills", "Travel"
    ],
    "Amount": [500, 200, 1200, 350, 1500, 250, 450, 800, 900, 300]
}

df = pd.DataFrame(data)

# Calculate summary
total_expenses = df["Amount"].sum()
average_expense = df["Amount"].mean()
highest_expense = df["Amount"].max()
category_expenses = df.groupby("Category")["Amount"].sum()

# Create bar chart
category_expenses.plot(kind="bar", title="Expenses by Category")
plt.xlabel("Category")
plt.ylabel("Amount (Rs.)")
plt.tight_layout()
plt.savefig("expense_chart.png")
plt.close()

# Export data to Excel
with pd.ExcelWriter("Personal_Expense_Dashboard.xlsx", engine="openpyxl") as writer:
    df.to_excel(writer, sheet_name="Expense Data", index=False)
    category_expenses.to_frame("Total Amount").to_excel(
        writer, sheet_name="Dashboard", startrow=6
    )


# Format Excel dashboard
wb = load_workbook("Personal_Expense_Dashboard.xlsx")
ws = wb["Dashboard"]

# Dashboard title
ws.merge_cells("A1:D1")
ws["A1"] = "PERSONAL EXPENSE DASHBOARD"
ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
ws["A1"].fill = PatternFill("solid", fgColor="305496")
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 30

# Summary values
ws["A3"] = "Total Expenses"
ws["B3"] = total_expenses
ws["A4"] = "Average Expense"
ws["B4"] = round(average_expense, 2)
ws["A5"] = "Highest Expense"
ws["B5"] = highest_expense

# Border style
thin_side = Side(style="thin", color="808080")
cell_border = Border(
    left=thin_side,
    right=thin_side,
    top=thin_side,
    bottom=thin_side
)

# Format summary boxes
for row in range(3, 6):
    for col in range(1, 3):
        cell = ws.cell(row=row, column=col)
        cell.border = cell_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.cell(row=row, column=1).font = Font(bold=True)
    ws.cell(row=row, column=1).fill = PatternFill(
        "solid", fgColor="D9EAF7"
    )
    ws.cell(row=row, column=2).font = Font(bold=True, color="305496")
    ws.cell(row=row, column=2).number_format = '"Rs. " #,##0.00'

# Format category table (rows 7 to 11)
for row in range(7, 12):
    for col in range(1, 3):
        cell = ws.cell(row=row, column=col)
        cell.border = cell_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

# Table heading
for cell in ws[7][0:2]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="4472C4")

# Amount formatting
for row in range(8, 12):
    ws.cell(row=row, column=2).number_format = '"Rs. " #,##0.00'

# Column widths
ws.column_dimensions["A"].width = 25
ws.column_dimensions["B"].width = 22
ws.column_dimensions["C"].width = 5
ws.column_dimensions["D"].width = 25

# Add chart
img = Image("expense_chart.png")
img.width = 520
img.height = 300
ws.add_image(img, "D3")

# Save Excel file
wb.save("Personal_Expense_Dashboard.xlsx")

print("Personal Expense Dashboard created successfully!")
print("Total Expenses: Rs.", total_expenses)
print("Excel file: Personal_Expense_Dashboard.xlsx")
